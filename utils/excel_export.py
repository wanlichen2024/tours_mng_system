import io
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from datetime import datetime
from utils.helpers import clean_display_value

def generate_itinerary_excel(selected_tour, hotels, attractions, meals):
    output = io.BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.title = "Itinerary"
    ws.append([f"Tour Code: {selected_tour}"])
    ws.append([f"Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"])
    ws.append([])

    if not hotels.empty:
        ws.append(["HOTEL BOOKINGS"])
        ws.append(["#", "City", "Supplier", "Check-in", "Check-out", "Rooms", "Product", "Reference", "Advance Notice (days)"])
        for i, (_, row) in enumerate(hotels.iterrows(), start=1):
            ws.append([
                i,
                clean_display_value(row.get('city')),
                clean_display_value(row.get('business_name')),
                clean_display_value(row.get('check_in_date')),
                clean_display_value(row.get('end_date')),
                clean_display_value(row.get('rooms')),
                clean_display_value(row.get('product_type')),
                clean_display_value(row.get('reference')),
                row.get('days_advance_notice', 32)
            ])
        ws.append([])

    if not attractions.empty:
        ws.append(["ATTRACTION BOOKINGS"])
        ws.append(["#", "City", "Supplier", "Date", "Time", "Pax", "Product", "Reference"])
        for i, (_, row) in enumerate(attractions.iterrows(), start=1):
            ws.append([
                i,
                clean_display_value(row.get('city')),
                clean_display_value(row.get('business_name')),
                clean_display_value(row.get('check_in_date')),
                clean_display_value(row.get('time')),
                clean_display_value(row.get('pax')),
                clean_display_value(row.get('product_type')),
                clean_display_value(row.get('reference'))
            ])
        ws.append([])

    if not meals.empty:
        ws.append(["RESTAURANT BOOKINGS"])
        ws.append(["#", "City", "Supplier", "Date", "Time", "Pax", "Product", "Reference"])
        for i, (_, row) in enumerate(meals.iterrows(), start=1):
            ws.append([
                i,
                clean_display_value(row.get('city')),
                clean_display_value(row.get('business_name')),
                clean_display_value(row.get('check_in_date')),
                clean_display_value(row.get('time')),
                clean_display_value(row.get('pax')),
                clean_display_value(row.get('product_type')),
                clean_display_value(row.get('reference'))
            ])

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            except:
                pass
        ws.column_dimensions[col_letter].width = min(max_len+2, 30)

    wb.save(output)
    output.seek(0)
    return output